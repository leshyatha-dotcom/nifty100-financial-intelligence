"""
Sprint 5 - PDF/Excel Tearsheet Generator
Generates a one-page PDF tearsheet and an Excel summary per company,
combining financial ratios, pros/cons, and NLP insights.
"""

import sqlite3
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
import openpyxl
from openpyxl.styles import Font, PatternFill


DB_PATH = "data/nifty100.db"
OUTPUT_DIR = "output/tearsheets"


def get_connection(db_path=DB_PATH):
    return sqlite3.connect(db_path)


def fetch_company_data(conn, company_id):
    company = conn.execute(
        "SELECT * FROM companies WHERE id = ?", (company_id,)
    ).fetchone()
    company_cols = [d[0] for d in conn.execute(
        "SELECT * FROM companies LIMIT 1"
    ).description]
    company_data = dict(zip(company_cols, company)) if company else {}

    ratios_row = conn.execute(
        "SELECT * FROM financial_ratios WHERE company_id = ? ORDER BY year DESC LIMIT 1",
        (company_id,)
    ).fetchone()
    ratio_cols = [d[0] for d in conn.execute(
        "SELECT * FROM financial_ratios LIMIT 1"
    ).description]
    ratios_data = dict(zip(ratio_cols, ratios_row)) if ratios_row else {}

    insight_row = conn.execute(
        "SELECT pros_summary, cons_summary FROM nlp_insights WHERE company_id = ?",
        (company_id,)
    ).fetchone()
    pros_summary, cons_summary = insight_row if insight_row else ("", "")

    return company_data, ratios_data, pros_summary, cons_summary


def generate_pdf_tearsheet(company_id, company_data, ratios_data, pros_summary, cons_summary, output_dir=OUTPUT_DIR):
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, f"{company_id}_tearsheet.pdf")

    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    name = company_data.get("company_name", company_id)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(2 * cm, y, f"{name} ({company_id})")
    y -= 1 * cm

    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Key Financial Ratios")
    y -= 0.6 * cm
    c.setFont("Helvetica", 10)

    shown = 0
    for key, val in ratios_data.items():
        if key in ("company_id", "id"):
            continue
        if val is None:
            continue
        c.drawString(2.2 * cm, y, f"{key}: {val}")
        y -= 0.5 * cm
        shown += 1
        if shown >= 10:
            break

    y -= 0.5 * cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Pros")
    y -= 0.6 * cm
    c.setFont("Helvetica", 9)
    for line in (pros_summary or "No data").split(" | ")[:6]:
        c.drawString(2.2 * cm, y, f"- {line[:95]}")
        y -= 0.45 * cm

    y -= 0.4 * cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Cons")
    y -= 0.6 * cm
    c.setFont("Helvetica", 9)
    for line in (cons_summary or "No data").split(" | ")[:6]:
        c.drawString(2.2 * cm, y, f"- {line[:95]}")
        y -= 0.45 * cm

    c.showPage()
    c.save()
    return file_path


def generate_excel_summary(all_rows, output_dir=OUTPUT_DIR):
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "all_companies_summary.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"

    headers = ["company_id", "company_name", "pros_summary", "cons_summary"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")

    for row in all_rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    wb.save(file_path)
    return file_path


def generate_all_tearsheets(db_path=DB_PATH, output_dir=OUTPUT_DIR):
    conn = get_connection(db_path)
    company_ids = [r[0] for r in conn.execute(
        "SELECT DISTINCT company_id FROM nlp_insights"
    ).fetchall()]

    excel_rows = []
    generated_pdfs = []

    for cid in company_ids:
        company_data, ratios_data, pros_summary, cons_summary = fetch_company_data(conn, cid)
        pdf_path = generate_pdf_tearsheet(cid, company_data, ratios_data, pros_summary, cons_summary, output_dir)
        generated_pdfs.append(pdf_path)
        excel_rows.append([
            cid,
            company_data.get("company_name", cid),
            pros_summary,
            cons_summary
        ])

    excel_path = generate_excel_summary(excel_rows, output_dir)
    conn.close()

    return generated_pdfs, excel_path


if __name__ == "__main__":
    pdfs, excel = generate_all_tearsheets()
    print(f"Generated {len(pdfs)} PDF tearsheets in output/tearsheets/")
    print(f"Generated Excel summary: {excel}")