"""
Sprint 3, Day 20: Generate output/peer_comparison.xlsx.

11 sheets (one per peer group). Each sheet: company_id, company_name, the
metric values, and a percentile-rank column per metric. Percentile cells are
colour-coded (green >=75th, yellow 25-75th, red <=25th). Benchmark company row
gets a gold/amber background. A summary row at the bottom shows the peer
group median for each metric.
"""

from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl.styles import PatternFill

BASE_DIR = Path(__file__).resolve().parents[2]
DB_PATH = BASE_DIR / "data" / "nifty100.db"

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOLD = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")

METRICS = [
    "return_on_equity_pct", "roce_pct", "net_profit_margin_pct", "debt_to_equity",
    "free_cash_flow_cr", "pat_cagr_5yr", "revenue_cagr_5yr", "eps_cagr_5yr",
    "interest_coverage", "asset_turnover",
]


def _pctile_fill(rank):
    if pd.isna(rank):
        return None
    if rank >= 0.75:
        return GREEN
    if rank <= 0.25:
        return RED
    return YELLOW


def build_peer_comparison():
    conn = sqlite3.connect(DB_PATH)
    peer_groups = pd.read_sql("SELECT * FROM peer_groups", conn)
    fr = pd.read_sql("SELECT * FROM financial_ratios", conn)
    percentiles = pd.read_sql("SELECT * FROM peer_percentiles", conn)
    companies = pd.read_sql("SELECT id, company_name FROM companies", conn)
    conn.close()

    fr_latest = fr.sort_values("year").groupby("company_id").tail(1)

    out_path = BASE_DIR / "output" / "peer_comparison.xlsx"
    out_path.parent.mkdir(exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for group_name, g in peer_groups.groupby("peer_group_name"):
            members = g["company_id"].tolist()
            benchmark = g[g["is_benchmark"] == 1]["company_id"].tolist()
            benchmark = benchmark[0] if benchmark else None

            sheet_df = fr_latest[fr_latest["company_id"].isin(members)].merge(
                companies, left_on="company_id", right_on="id", how="left"
            )
            cols = ["company_id", "company_name"] + [m for m in METRICS if m in sheet_df.columns]
            sheet_df = sheet_df[cols]

            # append percentile columns
            for metric in METRICS:
                pct_series = percentiles[
                    (percentiles["peer_group_name"] == group_name) & (percentiles["metric"] == metric)
                ].set_index("company_id")["percentile_rank"]
                sheet_df[f"{metric}_pctile"] = sheet_df["company_id"].map(pct_series)

            # summary row: peer group median
            summary = {"company_id": "MEDIAN", "company_name": ""}
            for metric in METRICS:
                if metric in sheet_df.columns:
                    summary[metric] = sheet_df[metric].median()
                summary[f"{metric}_pctile"] = None
            sheet_df = pd.concat([sheet_df, pd.DataFrame([summary])], ignore_index=True)

            sheet_name = group_name[:31]
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)

    return out_path, peer_groups


def apply_formatting(out_path, peer_groups):
    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    for group_name, g in peer_groups.groupby("peer_group_name"):
        sheet_name = group_name[:31]
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        benchmark = g[g["is_benchmark"] == 1]["company_id"].tolist()
        benchmark = benchmark[0] if benchmark else None

        id_col_idx = header.index("company_id") + 1

        for row_idx in range(2, ws.max_row + 1):
            company_id = ws.cell(row=row_idx, column=id_col_idx).value

            if company_id == benchmark:
                for col_idx in range(1, len(header) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = GOLD
                continue

            for col_name in header:
                if col_name.endswith("_pctile"):
                    col_idx = header.index(col_name) + 1
                    rank = ws.cell(row=row_idx, column=col_idx).value
                    fill = _pctile_fill(rank)
                    if fill:
                        ws.cell(row=row_idx, column=col_idx).fill = fill

    wb.save(out_path)


def main():
    out_path, peer_groups = build_peer_comparison()
    apply_formatting(out_path, peer_groups)
    print(f"peer_comparison.xlsx written: {out_path}")
    print(f"Sheets: {peer_groups['peer_group_name'].nunique()}")


if __name__ == "__main__":
    main()
